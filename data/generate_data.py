"""
Synthetic Dataset Generator for AI Resume Screening & ATS Score Predictor.
Generates realistic sample resumes in PDF, DOCX, and TXT formats across multiple
engineering disciplines, as well as production-grade job descriptions.
"""

import os
from typing import Dict, List

# Output directories
DATA_DIR = os.path.dirname(os.path.abspath(__file__))
RESUMES_DIR = os.path.join(DATA_DIR, "sample_resumes")
JDS_DIR = os.path.join(DATA_DIR, "sample_job_descriptions")

os.makedirs(RESUMES_DIR, exist_ok=True)
os.makedirs(JDS_DIR, exist_ok=True)

# Sample Candidate Data
CANDIDATES = [
    {
        "filename": "Alex_Chen_Senior_ML_Engineer",
        "name": "Alex Chen",
        "email": "alex.chen.ml@example.com",
        "phone": "(415) 892-4512",
        "linkedin": "linkedin.com/in/alexchen-ml",
        "github": "github.com/alexchen-ml",
        "education": "Master of Science in Computer Science | Stanford University (2018)",
        "experience_years": 6.0,
        "certs": "AWS Certified Machine Learning - Specialty, TensorFlow Developer Certificate",
        "summary": "Accomplished Senior Machine Learning Engineer with 6+ years of experience designing and deploying deep learning and NLP architectures in cloud-native production environments.",
        "skills": "Python, PyTorch, TensorFlow, Scikit-Learn, Transformers, Hugging Face, NLP, LLM, FastAPI, Docker, Kubernetes, AWS, SQL, CI/CD, Git",
        "experience": (
            "Senior Machine Learning Engineer | DeepLogic AI (2021 - Present)\n"
            "- Designed and implemented transformer-based RAG and search pipelines using PyTorch and Hugging Face, boosting retrieval NDCG@10 by 28%.\n"
            "- Containerized and deployed high-throughput inference endpoints with FastAPI, Docker, and Kubernetes on AWS EKS.\n"
            "- Built end-to-end continuous model training pipelines with GitHub Actions and MLflow.\n\n"
            "Machine Learning Engineer | DataVision Systems (2018 - 2021)\n"
            "- Developed production NLP classifiers and recommendation engines using Scikit-Learn and PyTorch.\n"
            "- Reduced inference latency by 45% via ONNX model quantization."
        ),
        "projects": (
            "Real-Time Semantic Document Search: Built a dense vector retrieval system using Sentence-Transformers, FastAPI, and ChromaDB.\n"
            "Autonomous Code Review Bot: Implemented an automated LLM reviewer using LangChain and GitHub Actions."
        )
    },
    {
        "filename": "Sarah_Johnson_FullStack_Developer",
        "name": "Sarah Johnson",
        "email": "sarah.johnson.dev@example.com",
        "phone": "(206) 451-9872",
        "linkedin": "linkedin.com/in/sarahj-dev",
        "github": "github.com/sarahj-dev",
        "education": "Bachelor of Science in Software Engineering | University of Washington (2019)",
        "experience_years": 5.0,
        "certs": "AWS Certified Developer - Associate",
        "summary": "Full Stack Software Engineer with 5 years of experience building modern responsive frontends with React & TypeScript, and high-performance backend APIs using Node.js and PostgreSQL.",
        "skills": "JavaScript, TypeScript, React, Node.js, Express.js, Python, PostgreSQL, Redis, Docker, Git, REST API, HTML5, CSS3, Tailwind CSS",
        "experience": (
            "Full Stack Software Engineer | CloudScale Inc. (2021 - Present)\n"
            "- Spearheaded frontend re-architecture to React 18, Next.js, and TypeScript, improving Core Web Vitals and load times by 40%.\n"
            "- Developed and optimized RESTful APIs and microservices with Node.js and PostgreSQL.\n"
            "- Implemented Redis caching layer reducing database read load by 55%.\n\n"
            "Frontend Developer | Nexus Labs (2019 - 2021)\n"
            "- Built interactive data visualization dashboards using React and Tailwind CSS."
        ),
        "projects": (
            "Collaborative Kanban Application: Real-time board built with React, WebSockets, and Node.js.\n"
            "E-Commerce Analytics Engine: Full-stack metrics dashboard powered by TypeScript and PostgreSQL."
        )
    },
    {
        "filename": "Michael_Patel_DevOps_Cloud_Architect",
        "name": "Michael Patel",
        "email": "michael.patel.infra@example.com",
        "phone": "(312) 762-3341",
        "linkedin": "linkedin.com/in/michaelpatel-cloud",
        "github": "github.com/mpatel-cloud",
        "education": "Bachelor of Science in Computer Information Systems | UIUC (2017)",
        "experience_years": 7.0,
        "certs": "AWS Certified Solutions Architect - Professional, Certified Kubernetes Administrator (CKA), HashiCorp Certified Terraform Associate",
        "summary": "Senior DevOps & Infrastructure Engineer with 7 years of hands-on experience provisioning multi-cloud environments, automating CI/CD workflows, and managing containerized Kubernetes clusters.",
        "skills": "AWS, Docker, Kubernetes, Terraform, Ansible, CI/CD, GitHub Actions, Jenkins, Prometheus, Grafana, Linux, Bash, Python, Helm",
        "experience": (
            "Staff Infrastructure Engineer | Enterprise Cloud Ops (2020 - Present)\n"
            "- Managed and maintained 15+ production Kubernetes clusters across AWS EKS using Helm and Terraform IaC.\n"
            "- Standardized CI/CD deployment workflows with GitHub Actions, accelerating release frequency by 3x.\n"
            "- Implemented comprehensive site reliability monitoring with Prometheus, Grafana, and Alertmanager, maintaining 99.99% uptime.\n\n"
            "Cloud Systems Administrator | Infotech Global (2017 - 2020)\n"
            "- Automated server provisioning and configuration using Ansible and Bash scripts across Linux environments."
        ),
        "projects": (
            "Multi-Region Kubernetes Failover: Automated Disaster Recovery setup with Terraform and Route 53.\n"
            "Zero-Trust GitOps Pipeline: Continuous deployment configuration with ArgoCD and Helm."
        )
    },
    {
        "filename": "Emily_Davis_Junior_Data_Analyst",
        "name": "Emily Davis",
        "email": "emily.davis.analyst@example.com",
        "phone": "(512) 634-1190",
        "linkedin": "linkedin.com/in/emilydavis-data",
        "github": "github.com/emilydavis-data",
        "education": "Bachelor of Science in Statistics | University of Texas at Austin (2023)",
        "experience_years": 1.5,
        "certs": "Google Data Analytics Professional Certificate",
        "summary": "Detail-oriented Junior Data Analyst with 1.5 years of experience performing statistical modeling, data cleaning, and creating executive dashboards in Tableau and Python.",
        "skills": "Python, SQL, Pandas, NumPy, Tableau, Excel, Data Cleaning, Statistical Analysis, A/B Testing",
        "experience": (
            "Junior Data Analyst | MetricFlow Analytics (2023 - Present)\n"
            "- Extracted and transformed monthly business metrics from PostgreSQL databases using advanced SQL queries.\n"
            "- Built automated reporting pipelines using Python and Pandas, saving 15 hours of manual reporting per week.\n"
            "- Created executive KPI dashboards in Tableau used by leadership team for quarterly reviews."
        ),
        "projects": (
            "Customer Churn Prediction: Performed exploratory data analysis and feature engineering on telecom dataset using Python and Pandas.\n"
            "Sales Cohort Analysis: Visualized user retention patterns with SQL and Seaborn."
        )
    }
]

# Sample Job Descriptions
JOB_DESCRIPTIONS = [
    {
        "filename": "Job_Senior_ML_Engineer.txt",
        "content": (
            "Job Title: Senior Machine Learning Engineer\n"
            "Department: AI & Machine Learning Research\n"
            "Experience Required: 5+ years of experience\n"
            "Education: Master's Degree or Bachelor's Degree in Computer Science, Data Science, or related field\n\n"
            "About the Role:\n"
            "We are seeking an experienced Senior Machine Learning Engineer to lead the architecture, training, and deployment of production NLP and large language model systems.\n\n"
            "Key Responsibilities:\n"
            "- Architect and deploy deep learning models for Natural Language Processing (NLP) and semantic retrieval.\n"
            "- Fine-tune state-of-the-art Transformer architectures using PyTorch and Hugging Face.\n"
            "- Deploy scalable inference microservices using FastAPI, Docker, and Kubernetes on AWS.\n"
            "- Design end-to-end continuous integration and deployment (CI/CD) pipelines for ML systems.\n"
            "- Collaborate across cross-functional teams to integrate generative AI models with core products.\n\n"
            "Required Qualifications & Skills:\n"
            "- 5+ years experience in Python and machine learning engineering.\n"
            "- Hands-on proficiency with PyTorch, Scikit-Learn, Transformers, and NLP libraries.\n"
            "- Strong experience with Docker, Kubernetes, and AWS cloud infrastructure.\n"
            "- Experience building RESTful APIs with FastAPI or Flask.\n"
            "- Solid foundation in SQL, Git, and automated CI/CD workflows."
        )
    },
    {
        "filename": "Job_FullStack_Developer.txt",
        "content": (
            "Job Title: Full Stack Web Developer\n"
            "Department: Product Engineering\n"
            "Experience Required: 4+ years of experience\n"
            "Education: Bachelor's Degree in Computer Science or Software Engineering\n\n"
            "About the Role:\n"
            "We are looking for a versatile Full Stack Developer to build elegant web applications and robust cloud microservices.\n\n"
            "Key Responsibilities:\n"
            "- Build modern, accessible user interfaces using React, TypeScript, and Tailwind CSS.\n"
            "- Design and implement backend REST APIs with Node.js and PostgreSQL.\n"
            "- Optimize database queries, indexing, and Redis caching layers.\n"
            "- Write clean, maintainable code with unit tests and Dockerized deployment workflows.\n\n"
            "Required Qualifications & Skills:\n"
            "- 4+ years of full-stack engineering experience.\n"
            "- Strong proficiency in JavaScript, TypeScript, React, and Node.js.\n"
            "- Hands-on experience with PostgreSQL, Redis, and REST APIs.\n"
            "- Familiarity with Docker, Git, and Agile methodologies."
        )
    },
    {
        "filename": "Job_DevOps_Cloud_Engineer.txt",
        "content": (
            "Job Title: Cloud DevOps & SRE Engineer\n"
            "Department: Cloud Platform & Infrastructure\n"
            "Experience Required: 4+ years of cloud engineering experience\n"
            "Education: Bachelor's Degree or equivalent experience\n\n"
            "About the Role:\n"
            "Join our infrastructure team to design, automate, and maintain mission-critical cloud platform services.\n\n"
            "Key Responsibilities:\n"
            "- Provision infrastructure as code (IaC) using Terraform across AWS.\n"
            "- Administer, scale, and secure production Kubernetes clusters.\n"
            "- Construct CI/CD pipelines with GitHub Actions and Docker.\n"
            "- Implement observability, logging, and metrics using Prometheus, Grafana, and Linux systems.\n\n"
            "Required Qualifications & Skills:\n"
            "- 4+ years experience in Cloud & DevOps engineering.\n"
            "- Deep knowledge of AWS, Docker, Kubernetes, and Terraform.\n"
            "- Strong Linux administration and Bash or Python scripting skills.\n"
            "- Experience with Prometheus, Grafana, and CI/CD tools."
        )
    }
]


def generate_txt_resume(cand: Dict, path: str):
    """Generates plain text resume."""
    text = (
        f"{cand['name']}\n"
        f"Email: {cand['email']} | Phone: {cand['phone']}\n"
        f"LinkedIn: {cand['linkedin']} | GitHub: {cand['github']}\n\n"
        f"PROFESSIONAL SUMMARY\n{cand['summary']}\n\n"
        f"TECHNICAL SKILLS\n{cand['skills']}\n\n"
        f"WORK EXPERIENCE\n{cand['experience']}\n\n"
        f"EDUCATION\n{cand['education']}\n\n"
        f"CERTIFICATIONS\n{cand['certs']}\n\n"
        f"PROJECTS\n{cand['projects']}\n"
    )
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)


def generate_pdf_resume(cand: Dict, path: str):
    """Generates PDF resume via ReportLab."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
        from reportlab.lib import colors

        doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()

        h1_style = ParagraphStyle('Name', parent=styles['Heading1'], fontSize=18, leading=22, textColor=colors.HexColor("#1E3A8A"))
        contact_style = ParagraphStyle('Contact', parent=styles['Normal'], fontSize=9, leading=12, textColor=colors.HexColor("#4B5563"))
        sec_style = ParagraphStyle('SecHead', parent=styles['Heading2'], fontSize=12, leading=16, textColor=colors.HexColor("#1F2937"), spaceBefore=10, spaceAfter=4)
        body_style = ParagraphStyle('Body', parent=styles['Normal'], fontSize=9.5, leading=13, textColor=colors.HexColor("#374151"))

        elements = [
            Paragraph(cand['name'], h1_style),
            Paragraph(f"{cand['email']} | {cand['phone']} | {cand['linkedin']} | {cand['github']}", contact_style),
            Spacer(1, 4),
            HRFlowable(width="100%", thickness=1, color=colors.HexColor("#CBD5E1"), spaceAfter=8),
            Paragraph("PROFESSIONAL SUMMARY", sec_style),
            Paragraph(cand['summary'], body_style),
            Paragraph("TECHNICAL SKILLS", sec_style),
            Paragraph(cand['skills'], body_style),
            Paragraph("WORK EXPERIENCE", sec_style),
        ]

        for line in cand['experience'].split("\n"):
            if line.strip():
                elements.append(Paragraph(line, body_style))
                elements.append(Spacer(1, 2))

        elements.extend([
            Paragraph("EDUCATION", sec_style),
            Paragraph(cand['education'], body_style),
            Paragraph("CERTIFICATIONS", sec_style),
            Paragraph(cand['certs'], body_style),
            Paragraph("NOTABLE PROJECTS", sec_style),
            Paragraph(cand['projects'], body_style),
        ])

        doc.build(elements)
    except Exception as e:
        # Fallback to TXT if ReportLab error
        print(f"Warning: Failed to generate PDF for {cand['name']}: {e}")
        generate_txt_resume(cand, path.replace(".pdf", ".txt"))


def generate_docx_resume(cand: Dict, path: str):
    """Generates DOCX resume via python-docx."""
    try:
        import docx
        doc = docx.Document()

        doc.add_heading(cand['name'], level=0)
        doc.add_paragraph(f"{cand['email']} | {cand['phone']} | {cand['linkedin']} | {cand['github']}")

        doc.add_heading("Professional Summary", level=1)
        doc.add_paragraph(cand['summary'])

        doc.add_heading("Technical Skills", level=1)
        doc.add_paragraph(cand['skills'])

        doc.add_heading("Work Experience", level=1)
        for line in cand['experience'].split("\n"):
            if line.strip():
                doc.add_paragraph(line)

        doc.add_heading("Education", level=1)
        doc.add_paragraph(cand['education'])

        doc.add_heading("Certifications", level=1)
        doc.add_paragraph(cand['certs'])

        doc.add_heading("Projects", level=1)
        doc.add_paragraph(cand['projects'])

        doc.save(path)
    except Exception as e:
        print(f"Warning: Failed to generate DOCX for {cand['name']}: {e}")
        generate_txt_resume(cand, path.replace(".docx", ".txt"))


def generate_all():
    """Generates all sample resumes and job descriptions."""
    print("Generating sample resumes and job descriptions...")

    # Generate JDs
    for jd in JOB_DESCRIPTIONS:
        jd_path = os.path.join(JDS_DIR, jd["filename"])
        with open(jd_path, "w", encoding="utf-8") as f:
            f.write(jd["content"])
        print(f"  Created JD: {jd_path}")

    # Generate Resumes in PDF, DOCX, and TXT
    for cand in CANDIDATES:
        base = cand["filename"]

        txt_p = os.path.join(RESUMES_DIR, f"{base}.txt")
        generate_txt_resume(cand, txt_p)
        print(f"  Created TXT Resume: {txt_p}")

        pdf_p = os.path.join(RESUMES_DIR, f"{base}.pdf")
        generate_pdf_resume(cand, pdf_p)
        print(f"  Created PDF Resume: {pdf_p}")

        docx_p = os.path.join(RESUMES_DIR, f"{base}.docx")
        generate_docx_resume(cand, docx_p)
        print(f"  Created DOCX Resume: {docx_p}")

    print("Synthetic dataset generation complete!")


if __name__ == "__main__":
    generate_all()
